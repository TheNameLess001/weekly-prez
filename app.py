import streamlit as st
import pandas as pd
import datetime
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# --- CONFIGURATION DE LA PAGE ---
st.set_page_config(page_title="Générateur WBR PPTX", page_icon="📊", layout="wide")

st.title("🚀 Générateur WBR LMD vers PowerPoint")
st.markdown("Importez vos fichiers (CSV pour Admin, XLSX pour le reste) et générez votre présentation PPTX.")

# --- FONCTIONS DE TRAITEMENT DES FICHIERS ---

def load_data(file):
    if file.name.endswith('.csv'):
        return pd.read_csv(file)
    else:
        return pd.read_excel(file)

@st.cache_data
def process_admin_file(file):
    df = pd.read_csv(file)
    df.columns = df.columns.str.strip().str.lower()
    if 'order day' in df.columns:
        df['order day'] = pd.to_datetime(df['order day'], errors='coerce')
    return df

@st.cache_data
def process_aux_files(files):
    df_sales_list = []
    df_cons_list = []
    
    for f in files:
        try:
            df = load_data(f)
            cols = df.columns.str.strip().str.lower()
            df.columns = cols
            
            if 'statut' in cols or 'segment' in cols:
                df_sales_list.append(df)
            elif "nom de l'établissement" in cols or "nom de l'etablissement" in cols:
                df_cons_list.append(df)
        except Exception as e:
            st.error(f"Erreur sur le fichier {f.name}: {e}")
            
    df_sales = pd.concat(df_sales_list, ignore_index=True) if df_sales_list else pd.DataFrame()
    df_cons = pd.concat(df_cons_list, ignore_index=True) if df_cons_list else pd.DataFrame()
    
    if not df_cons.empty and 'date' in df_cons.columns:
        df_cons['date'] = pd.to_datetime(df_cons['date'], errors='coerce')
        
    return df_sales, df_cons

def calculate_kpis(df):
    if df is None or df.empty:
        return {'orders': 0, 'gmv': 0, 'aov': 0, 'cancel_rate': 0, 'avg_time': 0, 'top_partners': [], 'cancel_reasons': []}
    
    df_delivered = df[df['status'] == 'Delivered']
    delivered_count = len(df_delivered)
    col_gmv = 'item total' if 'item total' in df.columns else 'payable amount'
    gmv = df_delivered[col_gmv].sum() if col_gmv in df.columns else 0
    aov = gmv / delivered_count if delivered_count > 0 else 0
    
    col_time = 'delivery time(m)'
    avg_time = df_delivered[col_time].mean() if col_time in df_delivered.columns else 0
    
    cancel_rate = (len(df[df['status'].astype(str).str.contains('Cancel', na=False, case=False)]) / len(df)) * 100 if len(df) > 0 else 0
    
    top_partners = []
    if 'restaurant name' in df.columns and not df_delivered.empty:
        top_partners = df_delivered.groupby('restaurant name')[col_gmv].sum().sort_values(ascending=False).head(10).reset_index().values.tolist()

    return {'orders': delivered_count, 'gmv': gmv, 'aov': aov, 'avg_time': avg_time, 'cancel_rate': cancel_rate, 'top_partners': top_partners}

# --- FONCTION GÉNÉRATION PPTX ---

def add_slide(prs, title_text, body_text_list=None, table_data=None):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    # Titre
    title = slide.shapes.title
    title.text = title_text
    
    # Corps de texte
    if body_text_list:
        tf = slide.placeholders[1].text_frame
        tf.text = body_text_list[0]
        for line in body_text_list[1:]:
            p = tf.add_paragraph()
            p.text = line
            p.level = 0

    # Tableau (si présent)
    if table_data:
        rows = len(table_data) + 1
        cols = len(table_data[0])
        left = Inches(0.5)
        top = Inches(2.5)
        width = Inches(9.0)
        height = Inches(0.8)
        
        table = slide.shapes.add_table(rows, cols, left, top, width, height).table
        # Header
        headers = ["Nom", "Valeur 1", "Valeur 2"] if cols == 3 else ["Item", "Valeur"]
        for i, h in enumerate(headers):
            table.cell(0, i).text = h
            
        for r_idx, row_data in enumerate(table_data):
            for c_idx, value in enumerate(row_data):
                table.cell(r_idx + 1, c_idx).text = str(value)

def generate_pptx(data):
    prs = Presentation()
    
    # Slide 1: Executive Summary
    add_slide(prs, "01. Executive Summary", [
        f"Commandes Livrées: {data['kpis']['orders']}",
        f"GMV Globale: {data['kpis']['gmv']:,.0f} MAD",
        f"AOV: {data['kpis']['aov']:.1f} MAD",
        "---",
        "Highlights:",
        data['in_high']
    ])
    
    # Slide 2: Operations
    add_slide(prs, "02. Operations & Delivery", [
        f"Temps moyen Click-to-Door: {data['kpis']['avg_time']:.0f} min",
        f"Taux d'annulation: {data['kpis']['cancel_rate']:.1f}%",
        f"Batching Rate: {data['in_batch']}",
        f"Flotte active: {data['in_fleet']} coursiers"
    ])
    
    # Slide 3: Sales & AM
    add_slide(prs, "03. Sales & Account Management", [
        f"Nouveaux Partenaires Signés: {data['in_signed']}",
        f"Onboardés: {data['in_onboarded']}",
        f"Top Deal: {data['in_topdeal']}"
    ])
    
    # Slide 4: Forecast Nouvelles Signatures
    add_slide(prs, "04. Forecast Nouvelles Signatures (Glovo vs Yassir)", None, data['forecast_table'])
    
    # Slide 5: Marketing & Growth
    add_slide(prs, "05. Performance Marketing", [
        f"Visites App: {data['in_visits']}",
        f"First Users (FTU): {data['in_ftu']}",
        f"ROI Promo: {data['in_gmv_promo']} GMV / {data['in_budget_promo']} Budget"
    ])
    
    # Slide 6: Blocages & Actions
    add_slide(prs, "06. Blocages & Difficultés", [
        "Blocages identifiés:",
        data['in_bloc']
    ])
    
    # Slide 7: Roadmap
    add_slide(prs, "07. Roadmap & Priorités S+1", [
        "Actions prioritaires:",
        data['in_road']
    ])
    
    binary_output = BytesIO()
    prs.save(binary_output)
    return binary_output.getvalue()

# --- INTERFACE STREAMLIT ---

admin_file = st.sidebar.file_uploader("Upload Admin Earnings (CSV)", type=['csv'])
aux_files = st.sidebar.file_uploader("Upload Sales/Consolider (XLSX)", type=['xlsx'], accept_multiple_files=True)

if admin_file:
    df_admin = process_admin_file(admin_file)
    df_sales, df_cons = process_aux_files(aux_files) if aux_files else (pd.DataFrame(), pd.DataFrame())
    
    # Inputs manuels (simplifiés pour l'exemple)
    st.subheader("Configuration de la présentation")
    in_high = st.text_area("Highlights", "Ex: Record de volume sur la zone Nord.")
    in_batch = st.number_input("Batching Rate", value=1.2)
    in_fleet = st.number_input("Coursiers actifs", value=450)
    in_signed = st.number_input("Signés", value=len(df_sales))
    in_onboarded = st.number_input("Onboardés", value=len(df_cons))
    in_topdeal = st.text_input("Top Deal", "McDonald's")
    in_visits = st.number_input("Visites App", value=15000)
    in_ftu = st.number_input("FTU", value=2000)
    in_budget_promo = st.number_input("Budget Promo", value=5000)
    in_gmv_promo = st.number_input("GMV Promo", value=50000)
    in_bloc = st.text_area("Blocages", "- Manque de livreurs le vendredi\n- Bug technique sur l'app")
    in_road = st.text_area("Roadmap S+1", "1. Lancement Zone B\n2. Campagne Promo Weekend")

    # Table Forecast
    st.write("Tableau de Forecast (Glovo vs Yassir)")
    df_f = st.data_editor(pd.DataFrame({"Restaurant": ["Resto A", "Resto B"], "Glovo": [100, 200], "Yassir": [80, 150]}), num_rows="dynamic")

    kpis = calculate_kpis(df_admin)

    data_for_pptx = {
        'kpis': kpis,
        'in_high': in_high,
        'in_batch': in_batch,
        'in_fleet': in_fleet,
        'in_signed': in_signed,
        'in_onboarded': in_onboarded,
        'in_topdeal': in_topdeal,
        'in_visits': in_visits,
        'in_ftu': in_ftu,
        'in_budget_promo': in_budget_promo,
        'in_gmv_promo': in_gmv_promo,
        'in_bloc': in_bloc,
        'in_road': in_road,
        'forecast_table': df_f.values.tolist()
    }

    if st.button("🎁 Générer le PowerPoint"):
        pptx_data = generate_pptx(data_for_pptx)
        st.download_button(
            label="📥 Télécharger le PPTX",
            data=pptx_data,
            file_name=f"WBR_LMD_{datetime.date.today()}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
