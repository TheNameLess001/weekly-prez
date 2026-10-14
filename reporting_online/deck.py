"""Native editable PowerPoint and optional local LibreOffice PDF conversion."""
from __future__ import annotations
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from PIL import Image
from engine import fmt, dumps

ROOT=Path(__file__).resolve().parent
DEFAULT=json.loads((ROOT/'assets/theme.json').read_text(encoding='utf-8'))

class Designer:
    def __init__(self,report,theme):
        self.r=report; self.t={**DEFAULT,**theme}; self.p=Presentation()
        self.p.slide_width=Inches(13.3333); self.p.slide_height=Inches(7.5)
        self.metrics={m['key']:m for m in report['metrics']}
    def color(self,key): return RGBColor.from_string(self.t.get(key,key).lstrip('#'))
    def box(self,s,x,y,w,h,color,rounded=False):
        sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb=self.color(color); sh.line.fill.background()
        if rounded: sh.adjustments[0]=.09
        return sh
    def text(self,s,x,y,w,h,text,size=18,color='text',bold=False):
        sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=sh.text_frame; tf.clear(); tf.word_wrap=True
        tf.margin_left=Pt(0); tf.margin_right=Pt(0); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0)
        for i,line in enumerate(str(text).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
            p.font.name=self.t['font']; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=self.color(color)
            p.space_after=Pt(6)
        return sh
    def logo(self,s,path,x,y,w,h):
        if not path or not Path(path).is_file(): return
        with Image.open(path) as im: iw,ih=im.size
        scale=min(w/iw,h/ih); ww,hh=iw*scale,ih*scale
        s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
    def new(self,title,subtitle='',dark=False):
        s=self.p.slides.add_slide(self.p.slide_layouts[6]); s.background.fill.solid()
        s.background.fill.fore_color.rgb=self.color('deep' if dark else 'background')
        self.text(s,.6,.4,10.9,.52,title,30,'white' if dark else 'text',True)
        if subtitle: self.text(s,.6,1.05,12,.48,subtitle,13,'white' if dark else 'muted')
        self.box(s,.6,7.05,12.1,.015,'primary')
        tag='DÉMONSTRATION FICTIVE' if self.r['demo'] else 'REPORTING PARTENAIRE'
        self.text(s,.6,7.18,11,.22,f"{tag}  /  {self.r['restaurant']['name'][:60]}  /  {self.r['restaurant']['id']}",10,'white' if dark else 'muted')
        self.text(s,12.2,7.15,.5,.25,str(len(self.p.slides)).zfill(2),11,'white' if dark else 'muted')
        return s
    def dates(self):
        a,b=self.r['current_dates']; c,d=self.r['previous_dates']
        return f"Actuel {a} au {b} ({self.r['current']['days']} j)   /   Réf. {c} au {d} ({self.r['previous']['days']} j)   /   Comparaison {self.r['comparison_basis']}"
    def card(self,s,key,x,y,w=3.86,h=1.6):
        m=self.metrics[key]; self.box(s,x,y,w,h,'white',True)
        self.text(s,x+.2,y+.15,w-.4,.43,m['label'],13,'muted')
        val=fmt(m['current'],m['kind'],self.t['currency'])
        self.text(s,x+.2,y+.56,w-.4,.44,val,26 if len(val)<18 else 21,'text',True)
        self.text(s,x+.2,y+1.1,w-.4,.24,f"Réf. {fmt(m['previous'],m['kind'],self.t['currency'])}",11,'muted')
        color='muted' if m['favorable'] is None else 'good' if m['favorable'] else 'bad'
        evol=m['delta_text'] if m['kind']=='rate' else m['growth']
        if m['previous']==0 and m['current'] is not None and m['current']!=0 and evol=='N/A':evol=m['delta_text']+' (réf. 0)'
        self.text(s,x+.2,y+1.35,w-.4,.23,evol,12,color,True)
    def grid(self,s,keys,y=1.8):
        for i,k in enumerate(keys): self.card(s,k,.6+(i%3)*4.08,y+(i//3)*1.9)
    def chart(self,s,kind,cats,series,x,y,w,h):
        if not cats or not series or any(any(v is None for v in vals) for _,vals in series):
            self.text(s,x+.25,y+.7,w-.5,1,'N/D — séries comparables indisponibles',20,'muted'); return
        data=CategoryChartData(); data.categories=cats
        for name,values in series: data.add_series(name,values)
        ch=s.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
        ch.font.name=self.t['font']; ch.font.size=Pt(11)
        ch.has_title=False
        from pptx.oxml.xmlchemy import OxmlElement
        auto_title=ch._chartSpace.chart.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}autoTitleDeleted')
        if auto_title is None:
            auto_title=OxmlElement('c:autoTitleDeleted'); ch._chartSpace.chart.insert(0,auto_title)
        auto_title.set('val','1')
        ch.has_legend=len(series)>1
        if ch.has_legend:
            ch.legend.position=XL_LEGEND_POSITION.BOTTOM
            ch.legend.font.size=Pt(12); ch.legend.font.name=self.t['font']
        ch.chart_style=10
        for i,ser in enumerate(ch.series):
            col=self.color('primary' if i%2==0 else 'muted')
            if kind==XL_CHART_TYPE.LINE:
                ser.format.line.color.rgb=col; ser.format.line.width=Pt(2.5)
            else: ser.format.fill.solid(); ser.format.fill.fore_color.rgb=col
        ch.category_axis.tick_labels.font.size=Pt(11); ch.category_axis.tick_labels.font.name=self.t['font']
        ch.value_axis.tick_labels.font.size=Pt(11); ch.value_axis.tick_labels.font.name=self.t['font']; ch.value_axis.minimum_scale=0
        ch.value_axis.has_major_gridlines=True
    def table(self,s,headers,rows,x,y,w,h,font=14,widths=None):
        tab=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h)).table
        if widths:
            for col,frac in zip(tab.columns,widths): col.width=Inches(w*frac)
        for i,row in enumerate([headers]+rows):
            for j,txt in enumerate(row):
                cell=tab.cell(i,j); cell.text=str(txt); cell.text_frame.word_wrap=True
                cell.fill.solid(); cell.fill.fore_color.rgb=self.color('deep' if i==0 else 'white')
                cell.margin_left=Inches(.11); cell.margin_right=Inches(.07); cell.margin_top=Inches(.055); cell.margin_bottom=Inches(.025)
                for p in cell.text_frame.paragraphs:
                    p.font.name=self.t['font']; p.font.size=Pt(font); p.font.bold=i==0; p.font.color.rgb=self.color('white' if i==0 else 'text')
        return tab
    def merch_empty(self,s,merch):
        self.box(s,.6,1.85,12.1,3.35,'white',True)
        zero=merch['current']['exact'] and merch['current']['totals']['quantity']==0
        self.text(s,.9,2.15,11.5,.6,'Aucune unité vendue sur la période certifiée' if zero else 'N/D — détail des ventes produits indisponible',26,'primary',True)
        body='Le fichier articles couvre la période complète et aucune ligne livrée n’est observée. Vérifier activation, catalogue et visibilité du restaurant.' if zero else 'Importer un fichier de lignes articles avec Restaurant ID, Order ID, Line ID, date, produit et quantité. Les catégories et le GMV ligne complètent l’analyse.\nLe classement porte uniquement sur les commandes livrées.'
        self.text(s,.9,3.0,11.3,1.35,body,19)
        notes=merch['current']['notes']+merch['previous']['notes']
        self.text(s,.9,5.6,11.4,.95,'\n'.join(notes[:2]) or 'Aucune vente produit observée sur les périodes sélectionnées.',16,'muted')
    def merch_slides(self):
        merch=self.r['merchandising']; cur=merch['current'];pre=merch['previous']
        products=[r for r in merch['products'] if r['current_qty'] is not None and r['current_qty']>0] if cur['rankable'] else []
        coverage=f"Articles : actuel {cur['covered_days']}/{cur['days']} j certifiés ; référence {pre['covered_days']}/{pre['days']} j certifiés."
        def short(s,n=46): return s if len(s)<=n else s[:n-1]+'…'
        s=self.new('Meilleures ventes',self.dates())
        if not products: self.merch_empty(s,merch)
        else:
            self.text(s,.7,1.7,11.8,.4,'Top 8 par quantité vendue sur les commandes livrées',20,bold=True)
            rows=[[f"{i+1:02d}  {short(r['name'])}",fmt(r['current_qty']),fmt(r['previous_qty']),fmt(r['current_gmv'],'money',self.t['currency']),fmt(r['share'],'rate'),r['growth']] for i,r in enumerate(products[:8])]
            self.table(s,['Produit','Unités','Réf.','GMV produit','Part unités','Évol. / jour'],rows,.6,2.35,12.1,3.65,13,[.34,.09,.09,.17,.13,.18])
            self.text(s,.7,6.25,11.8,.45,f"{fmt(cur['totals']['quantity'])} unités renseignées. {coverage}",13,'muted')
        s=self.new('Performance par catégorie',self.dates())
        categories=[r for r in merch['categories'] if r['current_qty'] is not None and r['current_qty']>0][:6] if cur['rankable'] else []
        if not categories: self.merch_empty(s,merch)
        else:
            self.text(s,.7,1.75,5.6,.4,'Mix des unités vendues',20,bold=True)
            self.chart(s,XL_CHART_TYPE.BAR_CLUSTERED,[short(r['name'],24) for r in categories],
                       [('Actuel',[r['current_qty'] for r in categories]),('Référence',[r['previous_qty'] for r in categories])],.6,2.3,5.7,3.6)
            rows=[[short(r['name'],26),fmt(r['share'],'rate'),fmt(r['current_gmv'],'money',self.t['currency']),f"{r['share_delta']*100:+.1f} pts".replace('.',',') if r['share_delta'] is not None else 'N/A',r['growth']] for r in categories]
            self.table(s,['Catégorie','Part unités','GMV produit','Écart mix','Évol. /j'],rows,6.7,2.35,6.0,3.5,11,[.31,.16,.16,.18,.19])
            self.text(s,.7,6.25,11.8,.47,'Parts calculées sur les unités produits renseignées. '+coverage,13,'muted')
        s=self.new('Top 3 des produits',self.dates())
        if not merch['top3']: self.merch_empty(s,merch)
        else:
            for i,r in enumerate(merch['top3']):
                x=.6+4.08*i;self.box(s,x,1.8,3.86,4.6,'white',True)
                self.text(s,x+.23,2.06,3.4,.6,f'#{i+1}',34,'primary',True)
                self.text(s,x+.23,2.95,3.4,.85,short(r['name'],50),20,bold=True)
                self.text(s,x+.23,3.87,3.4,.8,fmt(r['current_qty'])+' unités\n'+fmt(r['current_gmv'],'money',self.t['currency']),22)
                self.text(s,x+.23,5.0,3.4,.44,'Part des unités : '+fmt(r['share'],'rate'),14,'muted')
                color='good' if r['delta_per_day'] is not None and r['delta_per_day']>0 else 'bad' if r['delta_per_day'] is not None and r['delta_per_day']<0 else 'muted'
                self.text(s,x+.23,5.58,3.4,.55,'Évolution / jour : '+r['growth'],14,color,True)
            self.text(s,.7,6.6,11.8,.27,'Poids du Top 3 : '+fmt(merch['top3_share'],'rate')+f" des unités renseignées. Articles : {cur['covered_days']}/{cur['days']} j certifiés.",13,'muted')
        s=self.new('Hausses et Drop des produits','Évolutions en unités/jour. Des variations observées ne prouvent aucune causalité.')
        if not merch['reliable']:
            self.merch_empty(s,merch)
        else:
            for x,title,color,key in ((.6,'Top 3 en progression','good','risers'),(6.75,'Top 3 en baisse','bad','drops')):
                self.text(s,x+.1,1.85,5.8,.4,title,22,color,True)
                rr=merch[key]
                if rr:
                    rows=[[short(r['name'],30),f"{r['delta_per_day']:+.2f}".replace('.',','),r['growth']] for r in rr]
                    self.table(s,['Produit','Écart u./j','Évolution'],rows,x,2.55,5.95,2.8,13,[.5,.2,.3])
                else: self.text(s,x+.2,3.0,5.5,1.25,'Aucune '+('hausse' if key=='risers' else 'baisse')+' mesurable sur les données comparables.',19,'muted')
            text='Contrôler disponibilité, prix et visibilité des produits en baisse avant de proposer une relance.' if merch['drops'] else 'Suivre les produits en progression et vérifier leur disponibilité sur les prochains créneaux.'
            self.text(s,.7,5.85,11.9,.75,text,18)
            self.text(s,.7,6.65,11.8,.28,coverage,12,'muted')
    def build(self):
        name=self.r['restaurant']['name']; city=self.r['restaurant']['city']
        s=self.new('Performance partenaire',dark=True)
        fill=s.background.fill; fill.gradient(); fill.gradient_angle=35
        fill.gradient_stops[0].color.rgb=self.color('deep')
        deep=self.color('deep'); purple=self.color('primary')
        fill.gradient_stops[1].color.rgb=RGBColor(*(round(deep[i]*.8+purple[i]*.2) for i in range(3)))
        self.logo(s,self.t.get('yassir_logo'),10.7,.5,2,.65)
        if not self.t.get('yassir_logo'): self.text(s,10.7,.57,2,.4,'YASSIR',20,'white',True)
        self.text(s,.7,2.05,11.8,1.1,name,52 if len(name)<32 else 36,'white',True)
        self.text(s,.7,3.25,11.6,.55,city,23,'white')
        self.text(s,.7,4.3,11.6,1.2,self.dates().replace('   /   ','\n'),19,'white')
        self.logo(s,self.t.get('restaurant_logos',{}).get(self.r['restaurant']['id']),10.4,5.7,2.2,.8)
        self.text(s,.7,6.25,9,.35,'Données fictives. Aucun résultat réel.' if self.r['demo'] else 'Performance sur les périodes sélectionnées',15,'white')
        s=self.new('Synthèse exécutive',self.dates()); self.grid(s,['requests','delivered','sr','gmv','aov','users'],1.68)
        self.text(s,.7,5.7,11.9,1.12,'\n'.join(self.r['facts'][:3]),17)
        s=self.new('Performance commerciale',self.dates())
        self.text(s,.65,1.7,5.6,.4,'Volumes comparés',20,bold=True)
        self.chart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,['Reçues','Livrées'],
                   [('Actuel',[self.metrics[k]['current'] for k in ('requests','delivered')]),
                    ('Référence',[self.metrics[k]['previous'] for k in ('requests','delivered')])],.6,2.15,5.75,3.0)
        self.text(s,6.8,1.7,5.8,.4,'Évolution des livraisons',20,bold=True)
        trend=self.r['current']['trend']
        self.chart(s,XL_CHART_TYPE.LINE,[p['date'][5:] for p in trend],[('Actuel',[p['delivered'] for p in trend])],6.65,2.15,6.0,3.0)
        self.table(s,['Indicateur','Actuel','Référence','Évolution'],[
            [self.metrics[k]['label'],fmt(self.metrics[k]['current'],self.metrics[k]['kind']),fmt(self.metrics[k]['previous'],self.metrics[k]['kind']),self.metrics[k]['growth']]
            for k in ('gmv','delivered_per_day','gmv_per_day')],.7,5.35,11.9,1.3,13,[.34,.24,.24,.18])
        s=self.new('Clients et économie',self.dates())
        self.grid(s,['delivered_users','new_users' if self.metrics['new_users']['current'] is not None else 'observed_new_users','frequency','commission','promo_per_order','benefit'],1.7)
        self.text(s,.7,5.7,11.8,1.05,'Uniques : déduplication par client. Fréquence : livraisons / clients livrés.\nPremiers clients observés : historique des fichiers importés. Nouveaux clients globaux : référentiel de première commande requis.',16)
        s=self.new('Qualité opérationnelle',self.dates())
        self.grid(s,['accept_min','prep_min','wait_min','travel_min','ctd_min','cancel_rate'],1.7)
        m=self.metrics['sr']; cur=self.r['current']
        ns=', '.join(f"{k.replace('_min','')} n={int(v or 0)}" for k,v in cur['samples'].items() if k in ('accept_min','prep_min','ctd_min'))
        self.text(s,.7,5.72,11.9,.85,f"SR {fmt(m['current'],'rate')} ({m['delta_text']}). Annulations / commandes reçues.\nMoyennes des commandes livrées documentées : {ns or 'effectifs N/D'}.",16)
        reasons=sorted(cur['reasons'].items(),key=lambda v:-v[1])[:3]
        self.text(s,.7,6.58,11.9,.25,'Motifs : '+', '.join(f'{k[:40]} ({v})' for k,v in reasons) if reasons else 'Motifs d’annulation : N/D',12,'muted')
        s=self.new('Annulations et clients',self.dates())
        self.grid(s,['cancel_restaurant','cancel_driver','cancel_client','cancel_other','active_days','repeat_users'],1.7)
        other=self.metrics['cancel_other'];cancel=self.metrics['cancelled']
        self.text(s,.7,5.7,11.8,1.0,f"Annulations totales : {fmt(cancel['current'])} contre {fmt(cancel['previous'])}. Administration / autres : {fmt(other['current'])} contre {fmt(other['previous'])}. Acteur inconnu : {fmt(self.metrics['cancel_unknown']['current'])}.\nTous les taux d’annulation utilisent les commandes reçues au dénominateur. Premiers clients observés ≠ nouveaux clients globaux.",15)
        from engine import TIMES, TIME_LABELS
        s=self.new('Délais et effectifs',self.dates())
        rows=[]
        for t in TIMES:
            m=self.metrics[t];med=self.metrics[t+'_median']
            rows.append([TIME_LABELS[t],fmt(m['current'],'time'),fmt(m['previous'],'time'),m['growth'],fmt(med['current'],'time'),str(int(cur['samples'].get(t) or 0))+' / '+str(int(self.r['previous']['samples'].get(t) or 0))])
        self.table(s,['Durée','Moy. actuelle','Moy. réf.','Variation','Médiane act.','n act. / réf.'],rows,.6,1.8,12.1,3.3,14,[.24,.16,.16,.15,.16,.13])
        self.text(s,.7,5.45,11.8,1.2,'Préparation estimée = retrait − acceptation, attente éventuelle incluse. Préparation déclarée = prêt − acceptation. Attente après prêt = retrait − prêt.\nCes durées se recouvrent : ne pas les additionner. Les calculs portent sur les livraisons avec horodatages cohérents, de 0 à 1 440 minutes.',16)
        self.merch_slides()
        s=self.new('Plan d’action','Pistes à tester. Les variations observées ne prouvent aucune causalité.')
        for i,action in enumerate(self.r['actions']):
            y=1.75+i*1.67; self.box(s,.6,y,12.1,1.42,'white',True)
            self.text(s,.85,y+.13,3.1,.57,f"{i+1:02d}  {action['title']}",17,'primary',True)
            self.text(s,.85,y+.77,3.1,.55,action['fact'],12)
            self.text(s,4.3,y+.15,7.98,.7,action['action'],18)
            self.text(s,4.3,y+.97,7.98,.32,'Suivi : '+action['follow'],13,'muted')
        # Complete machine-verifiable KPI disclosure, including unavailable measures.
        ms=self.r['metrics']
        for offset in range(0,len(ms),10):
            s=self.new(f'Annexe KPI {offset//10+1}', 'Écarts et variations calculés avant arrondi. Actuel − référence. Taux en points, variation % sur base positive.')
            rows=[]
            for m in ms[offset:offset+10]:
                rows.append([m['label'],fmt(m['current'],m['kind']),fmt(m['previous'],m['kind']),m['delta_text'],m['growth']])
            self.table(s,['KPI','Actuel','Référence','Écart','Variation'],rows,.6,1.7,12.1,4.98,14,[.30,.19,.19,.16,.16])
        ca=self.r['current']; pr=self.r['previous']; defs=self.r['definitions']
        s=self.new('Définitions des indicateurs',self.dates())
        lines=[
            'SR = livraisons / reçues. AOV = GMV livré / livraisons. Fréquence = livraisons / clients livrés uniques.',
            'Récurrents = clients avec au moins 2 livraisons dans la période / clients livrés uniques. Déduplication des clients sur la période entière.',
            'Variation = (actuel − référence) / référence positive. Base zéro : nouvelle activité ou N/A. Base négative : écart absolu, pourcentage N/A.',
            'Les taux affichent leur écart en points de pourcentage. Les moyennes par jour divisent le total observé par les jours calendaires sélectionnés.',
            'Temps : livraisons avec une durée valide. Médianes calculées depuis les observations. Agrégats : moyennes pondérées par les effectifs.',
            'Acceptation = acceptation − création. Trajet = livraison − retrait. Click-to-door = durée source ou livraison − création.',
            'Jours actifs : jours avec une livraison observée. Disponibilité : heures disponibles / heures prévues, source requise.',
            'SLA : indicateur source ou click-to-door ≤ seuil utilisateur. Seuil : '+str(defs.get('sla_target_min') or 'non défini')+' min.'
        ]
        self.text(s,.7,1.7,11.8,5.1,'\n'.join(lines),18)
        s=self.new('Définitions financières',self.dates())
        for i,k in enumerate(('ca','commission','commission_base','promo','benefit')):
            y=1.7+i*.86
            self.text(s,.7,y,2.65,.6,k.replace('_',' ').capitalize(),17,'primary',True)
            self.text(s,3.5,y,9.0,.73,str(defs.get(k+'_definition') or 'Définition et source absentes : N/D'),16)
        self.text(s,.7,6.35,11.8,.45,'Convention monétaire : '+str(defs.get('tax_basis') or 'HT/TTC non qualifié dans la source.'),14,'muted')
        s=self.new('Sources et limites',self.dates())
        coverage=f"Jours avec observations : actuel {ca['observed_days']}/{ca['days']}, référence {pr['observed_days']}/{pr['days']}.\nJours certifiés exhaustifs : actuel {ca['covered_days']}/{ca['days']}, référence {pr['covered_days']}/{pr['days']}."
        self.text(s,.7,1.7,11.8,.9,coverage,20,'primary',True)
        notes=list(dict.fromkeys(ca['notes']+pr['notes']))
        self.text(s,.7,2.9,11.8,2.35,'\n'.join(notes[:4]) or 'Aucune limite de couverture détectée.',16)
        unavailable=[m for m in ms if m['current'] is None]
        text='\n'.join(m['label']+' : '+m['unavailable_reason'] for m in unavailable[:3])
        self.text(s,.7,5.25,11.8,1.0,text,14,'muted')
        sources=' ; '.join(self.r['sources'])
        self.text(s,.7,6.5,11.8,.35,'Sources : '+(sources[:190] or 'Portefeuille sans commandes')+' (détail dans les notes et le JSON)',12,'muted')
        # All exact provenance and definitions are also retained in editable notes.
        for slide in self.p.slides:
            slide.notes_slide.notes_text_frame.text=dumps({'sources':self.r['sources'],'definitions':defs,'limits':ca['notes']+pr['notes'],
                                                           'time_definitions':self.r['time_definitions'],'comparison_basis':self.r['comparison_basis'],
                                                           'samples_current':ca['samples'],'samples_previous':pr['samples'],'demo':self.r['demo'],
                                                           'product_sources':self.r['merchandising']['current']['sources']+self.r['merchandising']['previous']['sources'],
                                                           'product_definition':self.r['merchandising']['definition'],
                                                           'product_limits':self.r['merchandising']['current']['notes']+self.r['merchandising']['previous']['notes']})
        return self.p

def safe_name(value):
    s=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',str(value)).strip('. ')
    return (s[:65] or 'restaurant')

def pdf_converter():
    candidates=[shutil.which('soffice'),shutil.which('libreoffice'),os.environ.get('YASSIR_SOFFICE')]
    if os.name=='nt':
        candidates += [str(Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'LibreOffice/program/soffice.exe')]
    return next((c for c in candidates if c and Path(c).is_file()),None)

def convert_pdf(pptx_path,output):
    executable=pdf_converter()
    if not executable: return None,'PDF indisponible : installer LibreOffice localement (ou définir YASSIR_SOFFICE). PowerPoint reste disponible.'
    with tempfile.TemporaryDirectory(prefix='yassir_lo_') as profile:
        try:
            proc=subprocess.run([executable,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(output),str(pptx_path)],
                                capture_output=True,text=True,timeout=120)
            target=Path(output)/(Path(pptx_path).stem+'.pdf')
            if proc.returncode or not target.is_file(): return None,'Conversion PDF échouée : '+(proc.stderr or proc.stdout)[-600:]
            return target,None
        except (OSError,subprocess.TimeoutExpired) as e: return None,'Conversion PDF échouée : '+str(e)

def export_one(report,output,theme=None,pdf=False):
    out=Path(output); out.mkdir(parents=True,exist_ok=True)
    r=report['restaurant']; a,b=report['current_dates']; c,d=report['previous_dates']
    stem=f"{safe_name(r['name'])}_{safe_name(r['id'])}_{a}_{b}_vs_{c}_{d}"
    if report['demo']: stem='DEMO_FICTIF_'+stem
    target=out/(stem+'.pptx'); i=2
    while target.exists(): target=out/(stem+f'_{i}.pptx'); i+=1
    Designer(report,theme or {}).build().save(target)
    result={'pptx':str(target),'pdf':None,'warning':None}
    target.with_suffix('.json').write_text(dumps(report),encoding='utf-8')
    if pdf:
        f,w=convert_pdf(target,out); result.update(pdf=str(f) if f else None,warning=w)
    return result

def export_batch(ds,ids,current,previous,output,theme=None,pdf=False,progress=None):
    from engine import compare
    root=Path(output)/('lot_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    root.mkdir(parents=True,exist_ok=False); results=[]
    for n,rid in enumerate(ids):
        try: item={'id':rid,'ok':True,**export_one(compare(ds,rid,current,previous),root,theme,pdf)}
        except Exception as e: item={'id':rid,'ok':False,'error':str(e)}
        results.append(item)
        if progress: progress(n+1,len(ids),item)
    (root/'journal.json').write_text(dumps(results),encoding='utf-8')
    (root/'qualite.json').write_text(dumps(ds.quality),encoding='utf-8')
    archive=root.with_suffix('.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for f in root.iterdir(): z.write(f,f.name)
    return {'directory':str(root),'zip':str(archive),'results':results}
